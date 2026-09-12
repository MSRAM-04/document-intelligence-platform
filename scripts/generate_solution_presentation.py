from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)
ARCHITECTURE = DOCS / "system-architecture.png"
PRESENTATION = DOCS / "solution-presentation.pptx"

NAVY = RGBColor(20, 35, 52)
TEAL = RGBColor(0, 139, 139)
GOLD = RGBColor(225, 161, 72)
PALE = RGBColor(241, 246, 246)
INK = RGBColor(35, 47, 58)
WHITE = RGBColor(255, 255, 255)


def font(size, bold=False):
    try:
        return ImageFont.truetype("segoeui.ttf", size)
    except OSError:
        return ImageFont.load_default()


def make_architecture_image():
    image = Image.new("RGB", (1800, 1050), "#f2f6f5")
    draw = ImageDraw.Draw(image)
    title_font = font(52, True)
    body_font = font(28)
    small_font = font(24)
    draw.text((90, 55), "Ledger Lens | System Architecture", fill="#142334", font=title_font)
    draw.text((92, 125), "One deployable FastAPI service powers the dashboard and REST API.", fill="#3c5667", font=body_font)

    def box(x, y, w, h, title, lines, fill, outline="#0b7777"):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=22, fill=fill, outline=outline, width=4)
        draw.text((x + 24, y + 22), title, fill="#142334", font=font(30, True))
        for index, line in enumerate(lines):
            draw.text((x + 24, y + 78 + index * 36), line, fill="#304957", font=small_font)

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill="#d49a3a", width=7)
        draw.polygon([(x2, y2), (x2 - 18, y2 - 12), (x2 - 18, y2 + 12)], fill="#d49a3a")

    box(100, 260, 430, 250, "Web dashboard", ["HTML5 / CSS / JavaScript", "Upload, history, result views"], "#d9eeee")
    box(670, 215, 470, 340, "FastAPI REST service", ["Validation and error handling", "OCR and structured extraction", "Financial validation", "OpenAPI / Swagger"], "#cde5e3")
    box(1280, 260, 410, 250, "SQLite repository", ["JSON processing results", "Latest result by filename", "Dashboard list endpoint"], "#f9e5bd", outline="#c18b2d")
    box(115, 700, 360, 190, "Input validation", ["PDF / JPG / PNG", "20 MB, max 3 pages"], "#e4edf4", outline="#54758c")
    box(565, 700, 360, 190, "OCR layer", ["PyMuPDF page render", "EasyOCR text detection"], "#e4edf4", outline="#54758c")
    box(1015, 700, 360, 190, "Extraction + checks", ["Four document types", "Evidence and tolerances"], "#e4edf4", outline="#54758c")
    arrow(530, 385, 670, 385)
    arrow(1140, 385, 1280, 385)
    arrow(810, 555, 750, 700)
    arrow(750, 700, 790, 555)
    arrow(980, 555, 1190, 700)
    arrow(475, 795, 565, 795)
    arrow(925, 795, 1015, 795)
    draw.text((100, 970), "Render Web Service + persistent disk for deployed SQLite data", fill="#3c5667", font=body_font)
    image.save(ARCHITECTURE)


def add_text(slide, text, x, y, w, h, size=22, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align:
        paragraph.alignment = align
    return shape


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.7, 0.38, 12, 0.55, 28, NAVY, True)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.04), Inches(1.15), Inches(0.07)).fill.solid()
    accent = slide.shapes[-1]
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.7, 1.18, 12, 0.35, 13, RGBColor(80, 98, 109))


def add_bullets(slide, items, x=1.0, y=1.8, w=11.2, size=20, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.2))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(12)
        paragraph.text = "• " + paragraph.text
    return shape


def add_footer(slide, number):
    add_text(slide, f"Ledger Lens  |  Solution Presentation  |  {number}", 0.7, 7.05, 12, 0.22, 9, RGBColor(105, 122, 130))


def new_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    add_title(slide, title, subtitle)
    return slide


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "LEDGER LENS", 0.8, 1.25, 11.8, 0.7, 42, WHITE, True)
    add_text(slide, "Intelligent Document Intelligence Platform", 0.85, 2.05, 11.5, 0.8, 30, RGBColor(190, 236, 232), True)
    add_text(slide, "Extract, validate, evidence, and retrieve financial documents through one focused workflow.", 0.9, 3.15, 10.5, 0.6, 20, WHITE)
    add_text(slide, "Solution Presentation | AI Engineer Internship", 0.9, 5.85, 8, 0.35, 15, RGBColor(222, 188, 115))
    add_footer(slide, 1)

    slide = new_slide(prs, "The problem", "Financial documents arrive in inconsistent, semi-structured formats.")
    add_bullets(slide, ["Manual extraction is slow and difficult to audit.", "Scanned PDFs and image uploads require OCR before parsing.", "Extracted numbers need accounting consistency checks, not just text recognition.", "A reviewer needs both structured results and the evidence behind each field."], y=1.75)
    add_footer(slide, 2)

    slide = new_slide(prs, "Solution overview", "A compact pipeline turns an upload into a traceable processing result.")
    add_bullets(slide, ["Validate file type, size, readability, and page count.", "Render PDF pages and run EasyOCR over PDF, JPG, and PNG inputs.", "Parse four mandatory financial document types into structured JSON.", "Run formula checks and persist the latest result for retrieval and dashboard history."], y=1.75)
    add_footer(slide, 3)

    slide = new_slide(prs, "System architecture", "Frontend and API share one deployable FastAPI service.")
    slide.shapes.add_picture(str(ARCHITECTURE), Inches(0.7), Inches(1.55), width=Inches(11.95), height=Inches(5.25))
    add_footer(slide, 4)

    slide = new_slide(prs, "Processing pipeline", "Each stage has a clear responsibility and a testable output.")
    add_bullets(slide, ["Input: multipart upload plus document_type.", "Validation: PDF/JPG/PNG allowlist, empty/corrupt checks, 20 MB limit, three-page limit.", "OCR: PyMuPDF renders PDF pages; EasyOCR extracts spatial text.", "Extraction: document-specific fields, tables, evidence, page numbers, and confidence.", "Validation and storage: financial checks are returned in JSON before SQLite persistence."], y=1.65, size=18)
    add_footer(slide, 5)

    slide = new_slide(prs, "Supported documents", "The parser exposes consistent structured data across four financial workflows.")
    labels = [("Invoice", "Totals, tax, currency, line items"), ("Balance sheet", "Assets, liabilities, equity"), ("Profit & loss", "Revenue, costs, profit"), ("Cash flow", "Operating, investing, financing")]
    for index, (name, detail) in enumerate(labels):
        x = 0.85 + (index % 2) * 6.0
        y = 1.75 + (index // 2) * 2.0
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.35), Inches(1.35))
        shape.fill.solid(); shape.fill.fore_color.rgb = PALE; shape.line.color.rgb = TEAL
        add_text(slide, name, x + 0.25, y + 0.2, 4.8, 0.35, 21, TEAL, True)
        add_text(slide, detail, x + 0.25, y + 0.67, 4.8, 0.35, 15, INK)
    add_footer(slide, 6)

    slide = new_slide(prs, "API contract", "Simple REST endpoints cover upload, health, retrieval, and dashboard history.")
    add_bullets(slide, ["GET /api/v1/health -> service status.", "POST /api/v1/documents/process -> multipart file plus document_type.", "GET /api/v1/documents/{document_name} -> latest stored processing result.", "GET /api/v1/documents -> processed-document list for the dashboard.", "GET /docs -> generated Swagger/OpenAPI interface."], y=1.75, size=19)
    add_footer(slide, 7)

    slide = new_slide(prs, "Financial validation", "Numbers are checked as formulas and returned inside the JSON validation section.")
    add_bullets(slide, ["Invoice: subtotal + tax + shipping - discount = total.", "Balance sheet: assets approximately equal liabilities + equity.", "Profit and loss: revenue - cost of sales - operating expenses = net profit.", "Cash flow: operating + investing + financing = net change; opening + change = closing.", "Tolerance: $0.05 minimum absolute tolerance plus 1% relative tolerance; cash-flow translation tolerance is $5,000."], y=1.65, size=18)
    add_footer(slide, 8)

    slide = new_slide(prs, "Dashboard and persistence", "Processed documents remain discoverable after the upload request finishes.")
    add_bullets(slide, ["SQLite stores serialized processing results through a repository layer.", "The dashboard consumes the list endpoint and links to document results.", "GET-by-name returns the latest record for a document filename.", "Local path: data/documents.db; deployed path: /var/data/documents.db on the Render persistent disk."], y=1.75)
    add_footer(slide, 9)

    slide = new_slide(prs, "Deployment and responsible operation", "Render runs the same service that serves the browser dashboard and API.")
    add_bullets(slide, ["Platform: Render Web Service using render.yaml.", "Uvicorn binds to the platform PORT and serves / plus /api/v1.", "Configuration is supplied through environment variables; no secrets belong in Git.", "Persistent disk is required for durable SQLite history on the deployed service.", "Swagger URL and health endpoint provide quick smoke-test entry points."], y=1.75, size=19)
    add_footer(slide, 10)

    slide = new_slide(prs, "Submission verification", "Run these checks against the deployed URL before submitting the form.")
    add_bullets(slide, ["Open the frontend URL and upload a representative PDF, JPG, and PNG.", "Confirm GET /api/v1/health returns status ok.", "Confirm POST returns structured extracted_data and validation.checks.", "Confirm GET /api/v1/documents lists the processed record.", "Confirm GET /api/v1/documents/{name} returns the latest stored result.", "Replace the placeholders in README with the real Render and GitHub URLs."], y=1.6, size=18)
    add_footer(slide, 11)

    slide = new_slide(prs, "Limitations and next steps", "The current design is submission-ready and intentionally transparent about tradeoffs.")
    add_bullets(slide, ["OCR and first-load model initialization can be slow on small instances.", "Rule-based parsing is tuned to the supplied layouts and may need new templates for novel formats.", "SQLite is suitable for a single service instance, not high-volume multi-instance writes.", "Next: PostgreSQL/object storage, async OCR queue, authentication, rate limits, metrics, and deployment smoke tests."], y=1.75)
    add_footer(slide, 12)

    prs.save(PRESENTATION)


if __name__ == "__main__":
    make_architecture_image()
    build_presentation()
    print(f"Created {ARCHITECTURE}")
    print(f"Created {PRESENTATION}")
